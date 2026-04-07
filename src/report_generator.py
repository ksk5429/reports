"""Generate DOCX report in Korean with comprehensive restaurant analysis."""

import os
from datetime import datetime

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

from .models import Recommendation, SearchConstraints


def _set_cell_shading(cell, color_hex: str) -> None:
    shading = cell._element.get_or_add_tcPr()
    shading_elem = shading.makeelement(
        qn("w:shd"),
        {qn("w:fill"): color_hex, qn("w:val"): "clear"},
    )
    shading.append(shading_elem)


def _header_row(table, row_idx: int, color: str = "2F5496") -> None:
    for cell in table.rows[row_idx].cells:
        _set_cell_shading(cell, color)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True


def generate_report(
    recommendations: list[Recommendation],
    constraints: SearchConstraints,
    output_path: str,
    search_context: dict | None = None,
) -> str:
    """Generate comprehensive DOCX report in Korean."""
    doc = Document()
    ctx = search_context or {}

    # Page setup (A4)
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    for margin in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, margin, Cm(2.5))

    # ── 표지 ──
    title = doc.add_heading("맛집 추천 리포트", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(f"생성일: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(128, 128, 128)
    doc.add_paragraph("")

    # ── 1. 요약 ──
    doc.add_heading("1. 요약", level=1)
    doc.add_paragraph(
        f"본 리포트는 {ctx.get('date', '오늘')} 디너 데이트를 위한 "
        f"상위 {len(recommendations)}개 맛집 추천입니다. "
        f"DiningCode, CatchTable, 식신, 망고플레이트, TripAdvisor, 네이버, "
        f"인스타그램 등 10개 이상 플랫폼의 데이터를 종합 분석하였습니다."
    )

    # ── 2. 검색 조건 ──
    doc.add_heading("2. 검색 조건", level=1)
    items = [
        ("날짜", ctx.get("date", "N/A")),
        ("음식 종류", ctx.get("cuisine", "일식 (스키야키 / 이자카야)")),
        ("예산", ctx.get("budget", "1인 5~10만원")),
        ("인원", str(constraints.party_size) + "명"),
        ("분위기", ctx.get("atmosphere", "조용한 데이트 분위기")),
        ("주차", "필요 (여자친구 자차)" if constraints.parking_needed else "불필요"),
        ("웨이팅", "긴 대기 불가"),
    ]
    t = doc.add_table(rows=len(items) + 1, cols=2)
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.rows[0].cells[0].text = "항목"
    t.rows[0].cells[1].text = "값"
    _header_row(t, 0)
    for i, (k, v) in enumerate(items, 1):
        t.rows[i].cells[0].text = k
        t.rows[i].cells[1].text = v
    doc.add_paragraph("")

    # ── 3. 교통·동선 분석 ──
    doc.add_heading("3. 교통·동선 분석", level=1)
    logistics = ctx.get("logistics", [
        "여자친구: 인하대학교(인천) 18:00 출발, 자차 이동.",
        "퇴근시간(17~19시): 경인고속도로→서울 1.5~2시간 소요 예상.",
        "사용자: 개포레미안포레스트 출발.",
        "최적 합류 지점: 복정역 반경 20분 이내.",
        "예상 도착 시간: 19:30~20:00.",
        "추천 예약 시간: 19:30 또는 20:00.",
    ])
    for line in logistics:
        doc.add_paragraph(line, style="List Bullet")
    doc.add_paragraph("")

    # ── 4. 데이터 출처 ──
    doc.add_heading("4. 데이터 출처", level=1)
    platforms = [
        ("DiningCode", "평점, 메뉴, 리뷰"),
        ("CatchTable", "예약 가능 여부, 실시간 빈자리"),
        ("식신 (Siksinhot)", "사용자 평점, 큐레이션"),
        ("망고플레이트", "큐레이션 리뷰, 사진"),
        ("네이버 블로그/지도", "블로그 후기, 길찾기"),
        ("Google Maps", "해외 리뷰, 평점"),
        ("TripAdvisor / Trip.com", "여행자 리뷰"),
        ("Blue Ribbon Survey", "전문가 리뷰"),
        ("Instagram / SNS", "실시간 분위기, 해시태그"),
    ]
    pt = doc.add_table(rows=len(platforms) + 1, cols=2)
    pt.style = "Light Grid Accent 1"
    pt.alignment = WD_TABLE_ALIGNMENT.CENTER
    pt.rows[0].cells[0].text = "플랫폼"
    pt.rows[0].cells[1].text = "수집 데이터"
    _header_row(pt, 0)
    for i, (name, data) in enumerate(platforms, 1):
        pt.rows[i].cells[0].text = name
        pt.rows[i].cells[1].text = data
    doc.add_paragraph("")

    # ── 5. 상세 추천 ──
    doc.add_heading("5. 상세 추천", level=1)
    medals = {1: "🥇", 2: "🥈", 3: "🥉"}

    for rec in recommendations[:3]:
        r = rec.restaurant
        medal = medals.get(rec.rank, "")

        doc.add_heading(f"#{rec.rank} {medal} {r.name_ko} ({r.name})", level=2)

        info = [
            ("분류", f"{r.category} — {r.subcategory}"),
            ("주소", r.location.address),
            ("지역", r.location.district),
            ("전화", r.phone),
            ("영업시간", r.hours),
            ("브레이크타임", r.break_time),
            ("점심 가격", r.price.format_range("lunch")),
            ("저녁 가격", r.price.format_range("dinner")),
            ("주차", r.parking),
            ("예약", r.reservation_info),
            ("대기시간", r.wait_time_estimate),
            ("분위기", r.atmosphere),
            ("소음", r.noise_level),
            ("프라이빗 룸", "있음" if r.has_private_room else "없음"),
            ("노키즈존", "예" if r.no_kids_zone else "아니오"),
            ("데이트 점수", f"{r.date_friendly_score}/10"),
            ("종합 점수", f"{rec.composite_score}/10"),
        ]
        it = doc.add_table(rows=len(info), cols=2)
        it.style = "Light Grid Accent 1"
        it.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, (k, v) in enumerate(info):
            it.rows[i].cells[0].text = k
            it.rows[i].cells[1].text = str(v)
            for run in it.rows[i].cells[0].paragraphs[0].runs:
                run.bold = True
        doc.add_paragraph("")

        # 평점
        if r.ratings:
            doc.add_heading("플랫폼별 평점", level=3)
            rt = doc.add_table(rows=len(r.ratings) + 1, cols=5)
            rt.style = "Light Grid Accent 1"
            rt.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, h in enumerate(["출처", "점수", "맛", "서비스", "리뷰수"]):
                rt.rows[0].cells[j].text = h
            _header_row(rt, 0, "4472C4")
            for i, rating in enumerate(r.ratings, 1):
                rt.rows[i].cells[0].text = rating.source
                rt.rows[i].cells[1].text = f"{rating.score}/{rating.max_score}"
                rt.rows[i].cells[2].text = f"{rating.taste}" if rating.taste else "N/A"
                rt.rows[i].cells[3].text = f"{rating.service}" if rating.service else "N/A"
                rt.rows[i].cells[4].text = str(rating.review_count)
            doc.add_paragraph("")

        # 메뉴 하이라이트
        if r.highlights:
            doc.add_heading("메뉴 하이라이트", level=3)
            for h in r.highlights:
                doc.add_paragraph(h, style="List Bullet")

        # 장점
        if r.pros:
            doc.add_heading("장점", level=3)
            for p in r.pros:
                doc.add_paragraph(f"✓ {p}", style="List Bullet")

        # 주의사항
        if r.cons:
            doc.add_heading("주의사항", level=3)
            for c in r.cons:
                doc.add_paragraph(f"△ {c}", style="List Bullet")

        # 이동 시간
        if rec.travel_info_b:
            doc.add_heading("이동 시간 (여자친구: 인하대→)", level=3)
            doc.add_paragraph(f"예상 이동 시간: {rec.travel_info_b.estimated_duration_min}분")
            if rec.travel_info_b.traffic_notes:
                doc.add_paragraph(f"교통 참고: {rec.travel_info_b.traffic_notes}")

        if rec.travel_info_a:
            doc.add_heading("이동 시간 (본인: 개포→)", level=3)
            doc.add_paragraph(f"예상 이동 시간: {rec.travel_info_a.estimated_duration_min}분")

        # 출처
        if r.sources:
            doc.add_heading("출처", level=3)
            for s in r.sources:
                doc.add_paragraph(s, style="List Bullet")

        doc.add_page_break()

    # ── 6. 비교표 ──
    doc.add_heading("6. 비교표", level=1)
    headers = ["항목"] + [f"#{rec.rank} {rec.restaurant.name_ko}" for rec in recommendations[:3]]
    ct = doc.add_table(rows=9, cols=len(headers))
    ct.style = "Light Grid Accent 1"
    ct.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        ct.rows[0].cells[j].text = h
    _header_row(ct, 0)

    criteria = ["음식 퀄리티", "분위기", "가성비", "접근성", "예약 편의", "주차", "소음", "종합 점수"]
    for i, c in enumerate(criteria, 1):
        ct.rows[i].cells[0].text = c

    for col, rec in enumerate(recommendations[:3], 1):
        r = rec.restaurant
        vals = [
            f"{r.avg_rating * 5:.1f}/5",
            f"{r.date_friendly_score}/10",
            r.price.format_range("dinner"),
            f"{rec.travel_info_b.estimated_duration_min}분" if rec.travel_info_b else "N/A",
            (r.reservation_info[:30] if r.reservation_info else "N/A"),
            (r.parking[:30] if r.parking else "N/A"),
            r.noise_level,
            f"{rec.composite_score}/10",
        ]
        for row, val in enumerate(vals, 1):
            ct.rows[row].cells[col].text = val
    doc.add_paragraph("")

    # ── 7. 방법론 ──
    doc.add_heading("7. 분석 방법론", level=1)
    doc.add_paragraph(
        "본 추천은 10개 이상 플랫폼에서 수집한 데이터를 기반으로 합니다. "
        "각 맛집은 다음 가중치로 종합 점수를 산출합니다: "
        "음식 퀄리티(30%), 데이트 분위기(25%), 가성비(20%), 교통 편의성(25%). "
        "교통 분석은 하버사인 거리 + 서울 도심 보정계수를 적용했습니다."
    )

    # ── 8. 유의사항 ──
    doc.add_heading("8. 유의사항", level=1)
    doc.add_paragraph(
        "가격, 영업시간, 예약 가능 여부는 리포트 생성 시점 데이터이며 변경될 수 있습니다. "
        "방문 전 반드시 맛집에 직접 확인하시기 바랍니다. "
        "당일 예약은 캐치테이블 또는 전화로 확인하세요."
    )

    # Save
    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    doc.save(output_path)
    return output_path
